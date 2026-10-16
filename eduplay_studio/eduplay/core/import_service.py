"""
Import Service - Handles file imports (DOCX, PDF, XLSX)
"""

import os
import re
from typing import List, Dict, Optional
from pathlib import Path

from eduplay.core.i18n import I18n
from eduplay.core.path_resolver import PathResolver
import unicodedata


def _module_available(name: str) -> bool:
    """Kiem tra thu vien co cai dat hay khong MA KHONG import no.

    dung find_spec thay vi `try: import ...` de khoi dong app khong phai
    nap openpyxl (~0.9s + numpy) / docx / fitz ngay tu dau - chi nap khi
    nguoi dung that su bam "Nhap khau".
    """
    try:
        import importlib.util

        return importlib.util.find_spec(name) is not None
    except Exception:
        return False


# KHONG import o cap module: thu vien nang duoc nap lazy trong tung ham.
DocxDocument = None      # -> from docx import Document (khi can)
fitz = None              # -> import fitz (PyMuPDF, khi can)
load_workbook = None     # -> from openpyxl import load_workbook (khi can)

DOCX_AVAILABLE = _module_available("docx")
PDF_AVAILABLE = _module_available("fitz")
XLSX_AVAILABLE = _module_available("openpyxl")


def _lazy_docx_document():
    """Tra ve ham Document cua python-docx, nap lai khi can."""
    global DocxDocument
    if DocxDocument is None:
        from docx import Document as _Document

        DocxDocument = _Document
    return DocxDocument


def _lazy_fitz():
    """Tra ve module PyMuPDF, nap lai khi can."""
    global fitz
    if fitz is None:
        import fitz as _fitz

        fitz = _fitz
    return fitz


def _lazy_load_workbook():
    """Tra ve load_workbook cua openpyxl, nap lai khi can."""
    global load_workbook
    if load_workbook is None:
        from openpyxl import load_workbook as _load_workbook

        load_workbook = _load_workbook
    return load_workbook

# Bracket pairs [...] holding math expressions (e.g. [x+1]) are literal text,
# not blanks. A bracket is math when it holds an operator, so that answers
# like [100] or [Đại số] still count as blanks. Double brackets [[...]]
# always force a blank, even for math (teacher explicit).
_BRACKET_RE = re.compile(r"\[(.*?)\]")
_DOUBLE_BRACKET_RE = re.compile(r"\[\[(.*?)\]\]")
_COMBINED_RE = re.compile(r"\[\[(.*?)\]\]|\[(.*?)\]")
_MATH_OPS = frozenset("+-*/^=<>")

class ImportService:
    """Service for importing questions from various file formats"""
    
    def __init__(self):
        """Initialize import service"""
        self.templates_dir = Path(__file__).parent.parent.parent / "assets_bundle" / "templates"
    
    def get_available_templates(self) -> List[Dict]:
        """Get list of available sample templates"""
        templates = []
        
        if self.templates_dir.exists():
            for file_path in self.templates_dir.glob("*.txt"):
                template_name = file_path.stem.replace("_", " ").title()
                templates.append({
                    "name": template_name,
                    "file": str(file_path),
                    "description": I18n.t("import.template_desc_prefix") + template_name.lower()
                })
        
        return templates
    
    def load_template(self, template_name: str) -> List[Dict]:
        """Load questions from a template file"""
        template_path = self.templates_dir / f"{template_name}.txt"
        
        if not template_path.exists():
            # Try with underscores
            template_path = self.templates_dir / f"{template_name.replace(' ', '_').lower()}.txt"
        
        if template_path.exists():
            return self.import_from_txt(str(template_path))
        else:
            return []
    
    def _resolve_user_file(self, file_path: str) -> str:
        """Resolve user file path via PathResolver (OneDrive, file://, %20, long path)."""
        if not file_path:
            return str(file_path or "")
        try:
            # Try full resolver (file)
            p = PathResolver.resolve_user_file(str(file_path))
            if p and p.exists():
                return str(p)
            p2 = PathResolver.resolve_user_file_any(str(file_path))
            if p2 and p2.exists():
                return str(p2)
            # Fallback to normalized
            norm = PathResolver.normalize_user_path(str(file_path))
            if norm and Path(norm).exists():
                return norm
        except Exception:
            pass
        return str(file_path)

    def import_from_txt(self, file_path: str) -> List[Dict]:
        """Import questions from text file"""
        questions = []
        # OneDrive-aware resolve (handles file://, %20, 8.3, OneDrive sync)
        file_path = self._resolve_user_file(file_path)
        try:
            # Use utf-8-sig to handle BOM if present
            # Try resolver hydration first
            try:
                rp = PathResolver.resolve_user_file(file_path)
                if rp:
                    file_path = str(rp)
            except Exception:
                pass
            with open(file_path, 'r', encoding='utf-8-sig') as file:
                content = file.read()
            
            try:
                content = self._normalize_input(content)
            except Exception:
                pass
            
            # Parse questions based on format patterns
            # Try smart parse first as it handles blocks better
            smart_questions = self.parse_smart_format(content)
            if smart_questions:
                # Validate questions
                for q in smart_questions:
                    validated = self._validate_imported_question(q)
                    if validated:
                        questions.append(validated)
            
            # Legacy parsers are disabled to prevent duplicates and ensure consistent explanation extraction
            # questions.extend(self.parse_multiple_choice_questions(content))
            # questions.extend(self.parse_true_false_questions(content))
            # questions.extend(self.parse_fill_blank_questions(content))
            # questions.extend(self.parse_matching_questions(content))
            
            # If smart parser found nothing, we might try legacy? 
            # But we improved smart parser to handle everything.
            if not questions:
                 # Fallback just in case, but usually smart parser catches blocks
                 pass

            
        except Exception as e:
            print(f"Error importing from TXT: {e}")
        
        return questions
    
    def parse_smart_format(self, content: str) -> List[Dict]:
        """
        Smartly parse questions from text blocks.
        Handles:
        - Implicit options (A. B. - * or no prefix)
        - Implicit question types (MC, TF, FillBlank)
        - Implicit question text (first line(s))
        """
        try:
            content = self._normalize_input(content)
        except Exception:
            pass
        questions = []
        # Split by empty lines (2 or more newlines) OR by "Câu X:" headers OR by "X. Type" headers
        # This ensures we catch questions even if there are no empty lines between them.
        # Also supports headers like "2. Trắc nghiệm (Dạng khác)"
        type_keywords = r"(?:Trắc nghiệm|Đúng|Điền|Ghép|Tự luận|Multiple|True|Fill|Matching|Short)"
        pattern = r'(?:\n\s*\n)|(?=(?:^|\n)(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+\s*[:\.])|(?=(?:^|\n)\d+\.\s+.*' + type_keywords + r')'
        blocks = re.split(pattern, content, flags=re.MULTILINE)
        
        for block in blocks:
            lines = [line.strip() for line in block.split('\n') if line.strip()]
            if not lines:
                continue
            
            # Check for orphaned explanation (starts with "Giải thích:", "Explanation:", or "//")
            first_line_lower = lines[0].lower()
            is_explanation = (
                first_line_lower.startswith('giải thích:') or 
                first_line_lower.startswith('explanation:') or
                first_line_lower.startswith('//')
            )
            
            if is_explanation and questions:
                # Merge explanation to previous question
                expl_text = ""
                for line in lines:
                    if '//' in line:
                         parts = line.split('//', 1)
                         if len(parts) > 1:
                             content = parts[1].strip()
                             # Remove "Giải thích:" prefix if present in the comment
                             if content.lower().startswith('giải thích:'):
                                 content = content[11:].strip()
                             elif content.lower().startswith('explanation:'):
                                 content = content[12:].strip()
                             expl_text += content + "\n"
                    elif line.lower().startswith('giải thích:'):
                        expl_text += line[11:].strip() + "\n"
                    elif line.lower().startswith('explanation:'):
                        expl_text += line[12:].strip() + "\n"
                    else:
                        expl_text += line + "\n"
                
                prev_q = questions[-1]
                if prev_q.get("explanation"):
                    prev_q["explanation"] += "\n" + expl_text.strip()
                else:
                    prev_q["explanation"] = expl_text.strip()
                continue

            # Parse block
            q = self._parse_block(lines)
            if q:
                questions.append(q)
                
        return questions

    def _normalize_input(self, text: str) -> str:
        if not text:
            return ""
        t = text.replace('\r\n', '\n')
        t = t.replace('\u00A0', ' ')
        t = t.replace('\u200B', '')
        t = re.sub('[\u2000-\u200A\u202F\u205F]', ' ', t)
        return t

    def _remove_question_numbers(self, text: str) -> str:
        """Remove all question number prefixes from text"""
        if not text:
            return text
        # Remove patterns like "Câu 1:", "Question 1:", "câu 1", "questions 1", etc.
        text = re.sub(r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage|câu(?:\s+hỏi)?|question|pregunta|frage)\s+\d+\s*[:\.]?\s*', '', text, flags=re.IGNORECASE)
        # Also remove standalone number prefixes like "1.", "2)", "1:" at start of lines
        text = re.sub(r'^\s*\d+[\.\):]\s*', '', text, flags=re.MULTILINE)
        return text
 
    def _clean_header(self, lines: List[str]) -> List[str]:
        """Remove type header and question number prefixes from question lines"""
        if not lines:
            return lines

        # Strip empty lines from start
        # work on a copy to avoid side effects if original list is used elsewhere?
        # _parse_block passes slices mostly.
        lines = list(lines)

        while lines and not lines[0].strip():
            lines.pop(0)

        if not lines:
            return lines

        first_line = lines[0].strip()

        # Remove question number prefixes like "Câu 1:", "Question 1:", "câu 1", "questions 1", etc.
        # Pattern matches: Câu/Question/Pregunta/Frage followed by number and optional punctuation
        question_prefix_pattern = r'^(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage|câu(?:\s+hỏi)?|question|pregunta|frage)\s+\d+\s*[:\.]?\s*'
        if re.match(question_prefix_pattern, first_line, re.IGNORECASE):
            # Remove the prefix from the first line
            lines[0] = re.sub(question_prefix_pattern, '', first_line, flags=re.IGNORECASE).strip()
            # If the line becomes empty after removing prefix, remove it
            if not lines[0].strip():
                lines = lines[1:]

        # After removing question prefix, check if the remaining first line is a type header
        if lines:
            first_line = lines[0].strip()
            # Regex to match Type Headers (support all template headings across vi/en/fr/es/de)
            type_keywords = [
                r"Trắc nghiệm", r"Đúng\s*[\/-]\s*Sai", r"Điền khuyết", r"Điền vào chỗ trống", r"Ghép đôi", r"Ghép nối", r"Tự luận", r"Trả lời ngắn",
                r"Multiple Choice", r"True\s*[\/-]\s*False", r"Fill.*Blank", r"Matching", r"Short Answer",
                r"Choix multiple", r"Vrai\s*[\/-]\s*Faux", r"Remplir", r"Correspondance", r"Réponse courte",
                r"Opción múltiple", r"Verdadero\s*[\/-]\s*Falso", r"Llenar", r"Emparejar", r"Respuesta corta",
                r"Multiple.?Choice", r"Wahr\s*[\/-]\s*Falsch", r"Lückentext", r"Zuordnung", r"Kurzantwort"
            ]
            type_pattern = r'^(?:\d+[\.:]\s*)?.*(?:' + '|'.join(type_keywords) + r').*$'

            if re.match(type_pattern, first_line, re.IGNORECASE):
                lines = lines[1:]

        # Strip empty lines from start again
        while lines and not lines[0].strip():
            lines.pop(0)

        # Remove all question number prefixes from ALL lines (not just first line)
        # This handles cases where question numbers appear in the middle of text
        cleaned_lines = []
        for line in lines:
            # Remove patterns like "Câu 1:", "Question 1:", "câu 1", "questions 1", etc.
            cleaned_line = re.sub(r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage|câu(?:\s+hỏi)?|question|pregunta|frage)\s+\d+\s*[:\.]?\s*', '', line, flags=re.IGNORECASE)
            # Also remove standalone number prefixes like "1.", "2)", "1:" at start of lines
            cleaned_line = re.sub(r'^\s*\d+[\.\):]\s*', '', cleaned_line)
            cleaned_lines.append(cleaned_line)

        return cleaned_lines

    def _parse_block(self, lines: List[str]) -> Optional[Dict]:
        """Parse a single block of lines into a question"""
        if not lines:
            return None
            
        # Strategy:
        # 1. Identify explicit structure (Fill Blank with ==, or MC with Answer: line)
        # 2. Identify Options vs Question
        
        explanation = ""
        # Extract explanation first (lines starting with // or "Giải thích:" or "Explanation:", or inline //)
        filtered_lines = []
        for line in lines:
            # Check for inline comments //
            if re.search(r'(^|\s)//', line):
                parts = line.split('//', 1)
                content_part = parts[0].strip()
                expl_part = parts[1].strip()
                
                # Check if expl_part starts with "Giải thích:" or "Explanation:" (optional cleanup)
                if expl_part.lower().startswith('giải thích:'):
                    expl_part = expl_part[11:].strip()
                elif expl_part.lower().startswith('explanation:'):
                    expl_part = expl_part[12:].strip()
                
                explanation += expl_part + " "
                
                if content_part:
                    filtered_lines.append(content_part)
            
            # Check for full line legacy explanation markers
            elif line.lower().startswith('giải thích:'):
                explanation += line[11:].strip() + " "
            elif line.lower().startswith('explanation:'):
                explanation += line[12:].strip() + " "
            elif re.fullmatch(r'[\s=\-\*_~]+', line):
                # Skip decoration/separator lines like "====", "---", "***"
                continue
            else:
                filtered_lines.append(line)
        lines = filtered_lines
        
        if not lines:
            return None
        
        first_line_lower = lines[0].strip().lower()
        # Use regex to handle optional space before colon and case variations
        if (
            re.match(r'^\s*hướng dẫn\s*:', first_line_lower) or
            re.match(r'^\s*instructions\s*:', first_line_lower) or
            re.match(r'^\s*instrucciones\s*:', first_line_lower) or
            re.match(r'^\s*anleitung\s*:', first_line_lower) or
            ('hướng dẫn' in first_line_lower) or
            ('guidelines' in first_line_lower) or
            first_line_lower.startswith('instructions') or
            first_line_lower.startswith('instrucciones') or
            first_line_lower.startswith('anleitung')
        ):
            return None
        
        # Skip non-question header blocks like "Cấu trúc:" / "Structure:"
        if first_line_lower in ('cấu trúc:', 'structure:'):
            return None
        # Remove "Ví dụ:" / "Example:" headers but keep following lines for parsing
        if first_line_lower in ('ví dụ:', 'example:'):
            lines = lines[1:]
            if not lines:
                return None
        # Skip generic explanation/instruction sections
        if ('giải thích' in first_line_lower or 'explanation' in first_line_lower):
            return None

        # Check for Short Answer or Fill Blank (line starting with ==)
        has_q_header = any(re.match(r'^(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+\s*[:\.]', l, re.IGNORECASE) for l in lines)
        for i, line in enumerate(lines):
            if line.startswith('=='):
                tail = line[2:].strip()
                # Ignore lines that are pure decoration like "===="
                if not tail or re.fullmatch(r'[=\-\*_~]+', tail):
                    continue
                # If it looks like an option line (e.g. "A. ...", "1) ..."), treat it as MC correct marker, not short answer.
                try:
                    if re.match(r'^\s*([A-Fa-f]|\d{1,2})[\.\):]\s+.+', tail):
                        continue
                except Exception:
                    pass
                if has_q_header:
                    q_lines = self._clean_header(lines[:i])
                    question_text = "\n".join(q_lines).strip()
                    question_text = self._remove_question_numbers(question_text)
                    correct_answer = tail
                    
                    # Check if this is Fill Blank (has [blank] in question) or Short Answer (direct answer)
                    # If question has [blank] markers, it's fill_blank; math brackets
                    # (e.g. [x+1]) stay literal and don't count as blanks.
                    if re.search(r'\[.*?\]', question_text):
                        display_q, _blanks = self._split_brackets(question_text, [correct_answer])
                        if [b for b in (_blanks or []) if b]:
                            return {
                                "type": "fill_blank",
                                "question": display_q,
                                "correct_answer": correct_answer,
                                "correct_answers": [correct_answer],
                                "answers": [correct_answer],
                                "explanation": explanation.strip()
                            }
                    # Short answer: no brackets, or only literal math brackets
                    # (kept intact in the question text)
                    return {
                        "type": "short_answer",
                        "question": question_text,
                        "answers": [correct_answer],
                        "explanation": explanation.strip()
                    }

        # Check for Inline Fill Blank (e.g. "Water boils at [100] degrees.")
        # Only if we don't have explicit answer line (which might override)
        
        # Clean header first
        cleaned_lines = self._clean_header(lines)
        full_text = "\n".join(cleaned_lines)

        # Look for [answer] pattern
        # Brackets holding math (e.g. [x+5(78y+y^2)]) stay literal; only real
        # blanks count. Answers from '==' lines cross-check ambiguous brackets.
        _known_inline = []
        for _ln in lines:
            try:
                _s = str(_ln or "").strip()
            except Exception:
                continue
            if _s.startswith("=="):
                _t2 = re.sub(r"^\s*==\s*", "", _s)
                _t2 = re.sub(r"^([A-Fa-f]|\d{1,2})[\.\)]\s+", "", _t2).strip()
                if _t2:
                    _known_inline.append(_t2)
        fill_blank_matches = re.findall(r'\[(.*?)\]', full_text)
        if fill_blank_matches and has_q_header:
            display_q, _blanks_all = self._split_brackets(full_text, _known_inline)
            if not _blanks_all:
                # Only literal math brackets -> not a fill question, keep parsing below
                pass
            else:
                answers = [b for b in _blanks_all if b] or [""]
                question_text = self._remove_question_numbers(display_q)

                return {
                    "type": "fill_blank",
                    "question": question_text.strip(),
                    "answers": answers, # Use list for consistency
                    "correct_answers": answers, # For UI compatibility
                    "correct_answer": answers[0] if answers else "", # Fallback
                    "explanation": explanation.strip()
                }

        # Check for explicit "Answer:" line (Legacy support within block)
        explicit_answer_char = None
        legacy_answer_idx = -1
        
        for i, line in enumerate(lines):
            match = re.match(r'^\s*(?:[-*•]\s*)?(?:Đáp\s+án|Answer|Réponse|Respuesta|Antwort|Đáp\s+án\s+đúng|Correct\s+Answer)[:\s]\s*(.*)', line, re.IGNORECASE)
            if match:
                legacy_answer_idx = i
                explicit_answer_char = match.group(1).strip()
                try:
                    s_ans = explicit_answer_char
                    if ':' in s_ans:
                        s_ans = s_ans.split(':')[-1].strip()
                    m_letter = re.search(r'\b([A-Fa-f])\b', s_ans)
                    if m_letter:
                        explicit_answer_char = m_letter.group(1).upper()
                    else:
                        m_num = re.search(r'\b(\d{1,2})\b', s_ans)
                        if m_num:
                            explicit_answer_char = m_num.group(1)
                        else:
                            explicit_answer_char = s_ans.strip()
                except Exception:
                    pass
                # If explicit answer contains ==, strip it
                if explicit_answer_char.endswith('=='):
                    explicit_answer_char = explicit_answer_char[:-2].strip()
                elif explicit_answer_char.startswith('=='):
                     explicit_answer_char = explicit_answer_char[2:].strip()
                break
        
        # Remove Answer line from content processing if found
        if legacy_answer_idx != -1:
            lines.pop(legacy_answer_idx)

        # 2. Check for Answer Indicators
        # A valid question block MUST have either:
        # a) A line ending with '=='
        # b) An explicit "Answer:" line (explicit_answer_char)
        # c) Explicit True/False keywords as options (implied answer, though weak)
        # If none of these, it's likely instructions or headers -> Return None
        
        has_smart_answer = False
        smart_answer_line_idx = -1
        
        for i, line in enumerate(lines):
            s = line.strip()
            if s.endswith('==') or s.startswith('=='):
                has_smart_answer = True
                smart_answer_line_idx = i
                break
        
        if not has_smart_answer and not explicit_answer_char:
            has_matching = False
            match_count = 0
            for line in lines[1:]:
                s = line.strip()
                # Check for matching separators
                # We accept lines without bullet points if they look like pairs
                if ('->' in s or ' = ' in s or ' - ' in s) and '==' not in s:
                    match_count += 1
            has_matching = match_count >= 2
            
            # Check for True/False keywords which might be implicit TF
            # Only if lines look like options
            has_tf = False
            for line in lines[1:]: # Skip Q text
                clean = line.strip().lower()
                if clean in ["đúng", "sai", "true", "false", "vrai", "faux", "falsch", "wahr", "verdadero", "falso"]:
                    has_tf = True
                    break
            has_mc_prefix = False
            for line in lines[1:]:
                cl = line.strip()
                if re.match(r'^[A-Fa-f][\.\)]\s', cl) or cl.startswith('- ') or cl.startswith('* ') or re.match(r'^[1-9]\.\s', cl):
                    has_mc_prefix = True
                    break
            if not has_tf and not has_matching and not has_mc_prefix:
                return None
        
        # Identify Options
        # Heuristic:
        # - If any line ends with ==, it's an option (and correct).
        # - If any line starts with prefix (A., -), it's an option.
        # - If we found explicit_answer_char, we expect options.
        
        option_prefixes = (
            'A.', 'B.', 'C.', 'D.', 'E.', 'F.',
            'a.', 'b.', 'c.', 'd.', 'e.', 'f.',
            'A)', 'B)', 'C)', 'D)', 'E)', 'F)',
            'a)', 'b)', 'c)', 'd)', 'e)', 'f)',
            'A:', 'B:', 'C:', 'D:', 'E:', 'F:',
            'a:', 'b:', 'c:', 'd:', 'e:', 'f:',
            '- ', '* ', '• ', '– ',
            '1.', '2.', '3.', '4.',
            '1)', '2)', '3)', '4)'
        )
        
        # Determine split point between Question and Options
        split_idx = -1
        
        # Prefix families for smart grouping
        prefix_families = [
            ['A.', 'B.', 'C.', 'D.', 'E.', 'F.'],
            ['a.', 'b.', 'c.', 'd.', 'e.', 'f.'],
            ['A)', 'B)', 'C)', 'D)', 'E)', 'F)'],
            ['a)', 'b)', 'c)', 'd)', 'e)', 'f)'],
            ['A:', 'B:', 'C:', 'D:', 'E:', 'F:'],
            ['a:', 'b:', 'c:', 'd:', 'e:', 'f:'],
            ['1.', '2.', '3.', '4.'],
            ['1)', '2)', '3)', '4)'],
            ['- ', '* ', '• ', '– ']
        ]

        # Strategy A: Use Smart Answer ('==') location to find Option Block Start
        if smart_answer_line_idx != -1:
            # We found '=='. This line IS an option.
            # Scan backwards to find the start of the option block.
            
            # Identify prefix of the answer line (strip leading ==)
            ans_line = lines[smart_answer_line_idx].strip()
            # strip leading marker for prefix detection
            ans_line_noprefix = re.sub(r'^\s*==\s*', '', ans_line)
            ans_prefix = None
            current_family = None
            
            for p in option_prefixes:
                if ans_line_noprefix.startswith(p):
                    ans_prefix = p
                    break
            
            if ans_prefix:
                for family in prefix_families:
                    if ans_prefix in family:
                        current_family = family
                        break
            
            # Scan backwards
            curr_idx = smart_answer_line_idx
            first_option_idx = curr_idx
            
            for i in range(curr_idx - 1, -1, -1):
                line = lines[i].strip()
                if not line: continue
                
                is_option = False
                
                if re.match(r'^(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+\s*[:\.]', line, re.IGNORECASE):
                    break
                
                # Strip leading == for prefix detection
                line_noprefix = re.sub(r'^\s*==\s*', '', line)
                if ans_prefix:
                     # Check if line starts with a prefix from the SAME family
                     if current_family:
                         if any(line_noprefix.startswith(p) for p in current_family):
                             is_option = True
                     else:
                         # Fallback if prefix not in known families (shouldn't happen with standard list)
                         if any(line_noprefix.startswith(p) for p in option_prefixes):
                             is_option = True
                else:
                    first_option_idx = i
                    continue
                    
                if is_option:
                    first_option_idx = i
                elif ans_prefix:
                    # If we expect prefixes but this line has none or wrong family, it's likely the Question text
                    # E.g. "2. Header" (Family '1.') vs "A. Option" (Family 'A.')
                    break
                else:
                    break
            
            if ans_prefix is None and first_option_idx == 0:
                first_option_idx = 1
            split_idx = first_option_idx
            
        # Strategy B: Use First Prefix Match (Fallback if no smart answer or scanning failed)
        if split_idx == -1:
            # 1. Prioritize Letter Prefixes (A., B., etc.) to avoid matching numbered lists in question text
            letter_prefixes = [p for p in option_prefixes if p[0].lower() in 'abcdef']
            for i, line in enumerate(lines):
                if i == 0: continue
                clean_line = re.sub(r'^\s*==\s*', '', line.strip())
                if any(clean_line.startswith(p) for p in letter_prefixes):
                    split_idx = i
                    break
            
            # 2. If no letter prefixes, check other prefixes (Numbers, bullets)
            if split_idx == -1:
                other_prefixes = [p for p in option_prefixes if p[0].lower() not in 'abcdef']
                for i, line in enumerate(lines):
                    if i == 0: continue
                    clean_line = re.sub(r'^\s*==\s*', '', line.strip())
                    if any(clean_line.startswith(p) for p in other_prefixes):
                        split_idx = i
                        break
                
                # 3. Check for True/False keywords
                if split_idx == -1:
                    for i, line in enumerate(lines):
                        if i == 0: continue
                        clean_line = line.strip()
                        if clean_line.lower() in ["đúng", "sai", "true", "false", "vrai", "faux", "falsch", "wahr", "verdadero", "falso"]:
                            split_idx = i
                            break

            # 2. Check for Implicit Matching Pairs (lines with separators but no prefixes)
            if split_idx == -1:
                first_pair_idx = -1
                pair_count = 0
                for i, line in enumerate(lines):
                    if i == 0: continue # Skip question text
                    s = line.strip()
                    if ('->' in s or '=' in s or re.search(r'\s[-–—]\s', s)) and '==' not in s:
                        if first_pair_idx == -1:
                            first_pair_idx = i
                        pair_count += 1
                
                # If we found at least 2 pairs, assume this is a matching question
                if pair_count >= 2 and first_pair_idx != -1:
                    split_idx = first_pair_idx
        
        # 2. If no explicit marker found, but we have an explicit answer char (A, B..), 
        # try to find where options start based on A. B. (handled above) or just assume last N lines?
        # Or if we have implicit options (no prefix, no ==).
        if split_idx == -1:
            if explicit_answer_char:
                # Check if explicit answer is a TF keyword
                # If so, and we haven't found options, it's likely an Implicit True/False question
                # e.g. "Q: ... \n Answer: True"
                tf_keywords = ["đúng", "sai", "true", "false", "vrai", "faux", "falsch", "wahr", "verdadero", "falso"]
                if explicit_answer_char.lower() in tf_keywords:
                     # It is True/False!
                     is_true = explicit_answer_char.lower() in ["đúng", "true", "vrai", "verdadero", "wahr"]
                     question_text = self._remove_question_numbers("\n".join(lines).strip())
                     return {
                        "type": "true_false",
                        "question": question_text,
                        "correct_answer": is_true,
                        "explanation": explanation.strip()
                    }
                
                # If explicit answer is NOT TF, and we have no options, assume Fill Blank
                # But check if it looks like MC (A, B, C, D)
                if len(explicit_answer_char) == 1 and explicit_answer_char.upper() in "ABCD" and len(lines) > 1:
                     # Likely MC with implicit options?
                     pass
                else:
                     # Assume Fill Blank
                     question_text = self._remove_question_numbers("\n".join(lines).strip())
                     return {
                        "type": "fill_blank",
                        "question": question_text,
                        "correct_answers": [explicit_answer_char],
                        "answers": [explicit_answer_char],
                        "correct_answer": explicit_answer_char,
                        "explanation": explanation.strip()
                     }

            if len(lines) >= 3:
                # Heuristic: Line 0 is Question, Rest are Options
                split_idx = 1
            elif len(lines) == 2:
                # 2 lines. Q + Opt? Or Q + Answer (if explicit answer found)?
                # If explicit answer found, maybe lines are Q and Opts?
                # Case:
                # Q
                # True
                # Answer: True
                # -> split_idx should be 1.
                if explicit_answer_char:
                     split_idx = 1
        
        if split_idx == -1:
            # Can't find options. Maybe it's just a question text (Essay?) or incomplete.
            # If explicit_answer_char is present but no options found, maybe it's Short Answer/Fill Blank legacy?
            if explicit_answer_char:
                 question_text = self._remove_question_numbers("\n".join(lines).strip())
                 return {
                    "type": "fill_blank",
                    "question": question_text,
                    "correct_answer": explicit_answer_char,
                    "explanation": explanation.strip()
                }
            return None
            
        # Parse Options
        question_text_lines = lines[:split_idx]

        # Clean up Question Text: Remove "Type Headers"
        # Users often include "1. Multiple Choice", "2. True/False" etc. at the start.
        # We want to strip these lines so they don't appear in the actual question text.
        # But we must be careful not to strip valid question parts.
        # Heuristic: If the first line matches a known "Type Header" pattern, remove it.

        question_text_lines = self._clean_header(question_text_lines)

        # Strip empty lines from end of question text
        while question_text_lines and not question_text_lines[-1].strip():
            question_text_lines.pop()

        question_text = "\n".join(question_text_lines).strip()
        # Additional cleanup: remove any remaining question numbers from the final text
        question_text = self._remove_question_numbers(question_text)
        option_lines = lines[split_idx:]
        
        options = []
        for i, line in enumerate(option_lines):
            text = line.strip()
            is_correct = False
            
            # Leading correctness markers
            try:
                s0 = text.lstrip()
                if s0.startswith('=='):
                    is_correct = True
                    text = s0[2:].strip()
                elif s0.startswith('✓'):
                    is_correct = True
                    text = s0[1:].strip()
                elif s0.startswith('*') and len(s0) > 1:
                    is_correct = True
                    text = s0[1:].strip()
            except Exception:
                pass

            try:
                m_eq = re.search(r'\s*==\s*[.?!…]*\s*$', text)
                if m_eq:
                    is_correct = True
                    text = re.sub(r'\s*==\s*[.?!…]*\s*$', '', text).strip()
            except Exception:
                pass

            # Fallback: if '==' appears anywhere, treat it as a correctness marker
            try:
                if not is_correct and '==' in text:
                    is_correct = True
                    text = text.replace('==', ' ').strip()
            except Exception:
                pass
            
            matched_prefix = None
            for p in option_prefixes:
                if text.startswith(p):
                    matched_prefix = p
                    break
            
            if matched_prefix:
                text = text[len(matched_prefix):].strip()
                
                if explicit_answer_char:
                    clean_prefix = matched_prefix.strip(' .):')
                    if clean_prefix.lower() == explicit_answer_char.lower():
                        is_correct = True
            
            if explicit_answer_char and not is_correct:
                def _strip_accents(s: str) -> str:
                    try:
                        nf = unicodedata.normalize('NFD', s)
                        return ''.join(c for c in nf if not unicodedata.combining(c)).lower().strip()
                    except Exception:
                        return s.lower().strip()
                if text.lower().strip() == explicit_answer_char.lower().strip() or _strip_accents(text) == _strip_accents(explicit_answer_char):
                    is_correct = True
            
            try:
                m_annot = re.search(r'[\(\[\{]\s*(đúng|correct|true|vrai|wahr)\s*[\)\]\}]$', text, flags=re.IGNORECASE)
                m_tick = re.search(r'[\(\[\{]\s*✓\s*[\)\]\}]$', text)
                m_tick2 = re.search(r'✓\s*$', text)
                if m_annot or m_tick or m_tick2:
                    is_correct = True
                    text = re.sub(r'([\(\[\{]\s*(đúng|correct|true|vrai|wahr)\s*[\)\]\}])$', '', text, flags=re.IGNORECASE).strip()
                    text = re.sub(r'([\(\[\{]\s*✓\s*[\)\]\}])$', '', text).strip()
                    text = re.sub(r'✓\s*$', '', text).strip()
            except Exception:
                pass
            
            options.append({
                "text": text,
                "correct": is_correct
            })
            
        if not any(o['correct'] for o in options):
            expl_l = explanation.strip().lower()
            guess_idx = -1
            if expl_l:
                for i, o in enumerate(options):
                    t = o['text'].strip().lower()
                    if t and t in expl_l:
                        guess_idx = i
                        break
                if guess_idx == -1:
                    nums = []
                    for i, o in enumerate(options):
                        m = re.findall(r'\d+', o['text'])
                        if m:
                            try:
                                nums.append((i, int(m[-1])))
                            except:
                                pass
                    if nums:
                        target = None
                        m2 = re.findall(r'\d+', expl_l)
                        if m2:
                            try:
                                target = int(m2[-1])
                            except:
                                target = None
                        if target is not None:
                            for i, n in nums:
                                if n == target:
                                    guess_idx = i
                                    break
                        if guess_idx == -1:
                            nums.sort(key=lambda x: x[1], reverse=True)
                            guess_idx = nums[0][0]
            if guess_idx != -1 and 0 <= guess_idx < len(options):
                for i in range(len(options)):
                    options[i]['correct'] = (i == guess_idx)

        # Determine Type
        
        # Check if it is a Fill Blank Question (Smart Format)
        # Criteria: Only 1 option provided and it is marked correct (e.g. "Answer ==")
        # OR if it's explicitly marked as Fill Blank (via keywords) and has 1 option
        is_fill_blank_header = re.search(r'(?:Điền|Fill)', question_text, re.IGNORECASE)
        
        if (len(options) == 1 and options[0]['correct']) or (len(options) == 1 and is_fill_blank_header):
             return {
                "type": "fill_blank",
                "question": question_text,
                "correct_answers": [options[0]['text']],
                "answers": [options[0]['text']], # Legacy support
                "correct_answer": options[0]['text'], # UI compatibility
                "explanation": explanation.strip()
            }

        # Check if it is a Matching Question (Smart Format)
        # Criteria: Options contain '->' or '=' or ' - ', look like pairs
        
        # Relaxed matching criteria: If >= 50% of options have separators, treat as Matching
        # This handles cases where one line might be malformed or header-like
        arrow_count = sum(1 for o in options if '->' in o['text'])
        equal_count = sum(1 for o in options if ' = ' in o['text'])
        hyphen_count = sum(1 for o in options if re.search(r'\s[-–—]\s', o['text']))
        
        threshold = len(options) * 0.5
        has_arrow = arrow_count >= threshold
        has_equal = equal_count >= threshold
        has_hyphen = hyphen_count >= threshold
        
        if (has_arrow or has_equal or has_hyphen) and len(options) >= 2:
            pairs = []
            for o in options:
                parts = []
                if '->' in o['text']:
                    parts = o['text'].split('->', 1)
                elif '=' in o['text']:
                    parts = o['text'].split('=', 1)
                elif re.search(r'\s[-–—]\s', o['text']):
                    if ' - ' in o['text']:
                        parts = o['text'].split(' - ', 1)
                    elif ' – ' in o['text']:
                        parts = o['text'].split(' – ', 1)
                    elif ' — ' in o['text']:
                        parts = o['text'].split(' — ', 1)
                else:
                    continue
                    
                if len(parts) >= 2:
                    pairs.append({
                        "left": parts[0].strip(),
                        "right": parts[1].strip()
                    })
            
            if pairs:
                return {
                    "type": "matching",
                    "question": question_text,
                    "pairs": pairs,
                    "explanation": explanation.strip()
                }

        # Check if it is True/False
        # Criteria: 2 options, one True, one False (keywords)
        # Or explicit True/False keywords found
        tf_keywords = ["đúng", "sai", "true", "false", "vrai", "faux", "falsch", "wahr", "verdadero", "falso"]
        is_tf = False
        if len(options) == 2:
            opt_texts = [o['text'].lower() for o in options]
            if all(t in tf_keywords for t in opt_texts):
                is_tf = True
        
        if is_tf:
            # Determine correct answer (Boolean)
            correct_bool = True
            # Find which one is marked correct
            correct_opt = next((o for o in options if o['correct']), None)
            
            if correct_opt:
                if correct_opt['text'].lower() in ["sai", "false", "faux", "falsch", "falso"]:
                    correct_bool = False
            else:
                # If no marker, maybe check explicit_answer_char?
                if explicit_answer_char:
                    if explicit_answer_char.lower() in ["sai", "false", "faux", "falsch", "falso", "b"]:
                        correct_bool = False
            
            return {
                "type": "true_false",
                "question": question_text,
                "correct_answer": correct_bool,
                "explanation": explanation.strip()
            }

        # Default: Multiple Choice
        # We need to format options as Strings ("Content", "Content") - No prefix!
        # and set correct_answer to the Index (0, 1, 2, 3)
        
        formatted_options = []
        letters = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']
        correct_index = 0
        
        for i, opt in enumerate(options):
            # We used to add prefix here, but UI might prefer clean text
            # formatted_options.append(f"{letters[i] if i < len(letters) else str(i+1)}. {opt['text']}")
            formatted_options.append(opt['text'])
            
            if opt['correct']:
                correct_index = i
        
        # If explicit answer char was provided (e.g. "Answer: B"), trust it if it matches range
        if explicit_answer_char:
            clean_char = explicit_answer_char.strip(" .)")
            if len(clean_char) == 1 and clean_char.upper() in letters[:len(options)]:
                correct_letter = clean_char.upper()
                correct_index = letters.index(correct_letter)
            elif clean_char.isdigit():
                try:
                    num = int(clean_char)
                    if 1 <= num <= len(options):
                        correct_index = num - 1
                except Exception:
                    pass
        
        return {
            "type": "multiple_choice",
            "question": question_text,
            "options": formatted_options,
            "correct_answer": correct_index, # UI expects index
            "explanation": explanation.strip()
        }

    def parse_multiple_choice_questions(self, content: str) -> List[Dict]:
        """Parse multiple choice questions from text"""
        questions = []
        
        # Flexible pattern:
        # Capture from "Question X" to "Answer: Y"
        # We capture the body (text + options) and the answer
        pattern = r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+(\d+)\s*[:\.]\s*(.+?)\s*(?:Đáp\s+án|Answer|Réponse|Respuesta|Antwort|Đáp\s+án\s+đúng|Correct\s+Answer):\s*([A-D])'
        
        # Split content into blocks roughly to avoid greedy matching issues across questions?
        # Actually, findall with non-greedy (.+?) usually works well for sequential blocks
        matches = re.findall(pattern, content, re.MULTILINE | re.DOTALL)
        
        for match in matches:
            question_num, body, correct_answer = match
            
            # Now extract options from the body
            # Look for lines starting with A., B., C., D. (or A), B))
            # We assume options are at the end of the body
            
            # Regex for options: A. text
            opt_matches = list(re.finditer(r'(?:^|\n)\s*([A-D])[\.\)]\s*(.+?)(?=(?:\n\s*[A-D][\.\)])|$)', body, re.DOTALL))
            
            if not opt_matches:
                continue
                
            # Question text is everything before the first option
            first_opt_start = opt_matches[0].start()
            question_text = body[:first_opt_start].strip()
            
            options = []
            valid_options = True
            
            for m in opt_matches:
                label = m.group(1)
                text = m.group(2).strip()
                options.append({
                    "text": text,
                    "correct": label.upper() == correct_answer.upper()
                })
            
            if len(options) >= 2:
                # Extract explanation (look ahead after the match in original content)
                # This is hard with findall. 
                # Alternative: The "Answer" line was consumed. We check what's after.
                # But findall doesn't give us position.
                # We'll use a simpler explanation extraction or skip it for now.
                # Let's try to find explanation in the main loop if we iterate manually.
                
                # For now, let's just append the question.
                questions.append({
                    "type": "multiple_choice",
                    "question": question_text,
                    "options": options,
                    "explanation": "" 
                })
        
        return questions
    
    def parse_true_false_questions(self, content: str) -> List[Dict]:
        """Parse true/false questions from text"""
        questions = []
        
        # Pattern for true/false questions (supports VI/EN/FR/ES/DE)
        # Use negative lookahead to prevent matching across questions
        pattern = r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+(\d+)\s*[:\.]\s*((?:(?!\n\s*(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+).)+?)\n\s*(?:Đáp\s+án|Answer|Réponse|Respuesta|Antwort):\s*(Đúng|Sai|True|False|Vrai|Faux|Verdadero|Falso|Wahr|Falsch)'
        
        matches = re.findall(pattern, content, re.MULTILINE | re.DOTALL)
        
        for match in matches:
            question_num, question_text, correct_answer = match
            
            is_true = correct_answer.lower() in ("đúng", "true", "vrai", "verdadero", "wahr")
            
            questions.append({
                "type": "true_false",
                "question": question_text.strip(),
                "correct_answer": is_true,
                "explanation": ""
            })
        
        return questions
    
    def parse_fill_blank_questions(self, content: str) -> List[Dict]:
        """Parse fill-in-the-blank questions from text"""
        questions = []
        
        # Pattern for fill blank questions (supports VI/EN/FR/ES/DE)
        # Use lookahead to stop at next question or end of string
        pattern = r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+(\d+)\s*[:\.]\s*((?:(?!\n\s*(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+).)+?)\n\s*(?:Đáp\s+án|Answer|Réponse|Respuesta|Antwort):\s*(.+?)(?=\n\s*(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)|$)'
        
        matches = re.findall(pattern, content, re.MULTILINE | re.DOTALL)
        
        for match in matches:
            question_num, question_text, correct_answer = match
            answer_text = correct_answer.strip()
            
            # Skip if it looks like True/False
            if answer_text.lower() in ("đúng", "true", "vrai", "verdadero", "wahr", "sai", "false", "faux", "falso", "falsch"):
                continue
                
            # Skip if it looks like Multiple Choice (Single letter A-D and body has options)
            if len(answer_text) == 1 and answer_text.upper() in "ABCD":
                # Check if body contains options like "A. " or "A) "
                if re.search(r'(?:^|\n)\s*[A-D][\.\)]', question_text):
                    continue
            
            questions.append({
                "type": "fill_blank",
                "question": question_text.strip(),
                "answers": [answer_text],
                "case_sensitive": False,
                "explanation": ""
            })
        
        return questions
    
    def parse_matching_questions(self, content: str) -> List[Dict]:
        """Parse matching questions from text"""
        questions = []
        
        # Look for matching blocks
        # Pattern: Question X: ... \n Column A - Column B \n ...
        pattern = r'(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+(\d+)\s*[:\.]\s*((?:(?!\n\s*(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)\s+\d+).)+?)\n\s*(?:Cột A|Column A|Colonne A|Columna A|Spalte A)\s*-\s*(?:Cột B|Column B|Colonne B|Columna B|Spalte B)\s*\n(.+?)(?=\n\s*(?:Câu(?:\s+hỏi)?|Question|Pregunta|Frage)|$)'
        
        matches = re.findall(pattern, content, re.MULTILINE | re.DOTALL)
        
        for match in matches:
            question_num, question_text, pairs_text = match
            
            pairs = []
            for line in pairs_text.strip().split('\n'):
                line = line.strip()
                if not line: continue
                
                # Remove leading dash or bullet if present
                if line.startswith('- '):
                    line = line[2:]
                elif line.startswith('• '):
                    line = line[2:]
                elif line.startswith('-'): # Handle dash without space
                    line = line[1:].strip()
                
                # Split by -> or -
                if '->' in line:
                    parts = line.split('->', 1)
                elif '-' in line:
                    parts = line.split('-', 1)
                else:
                    continue

                if len(parts) == 2:
                    pairs.append({
                        "left": parts[0].strip(),
                        "right": parts[1].strip()
                    })
            
            if pairs:
                questions.append({
                    "type": "matching",
                    "question": question_text.strip(),
                    "pairs": pairs,
                    "explanation": ""
                })
        
        return questions
    
    def import_from_file(self, file_path: str) -> List[Dict]:
        """Import questions from file based on extension"""
        # OneDrive-aware resolve before existence check
        resolved = self._resolve_user_file(file_path)
        # Also try full resolver for cloud hydration
        try:
            pr = PathResolver.resolve_user_file(resolved)
            if pr and pr.exists():
                resolved = str(pr)
        except Exception:
            pass
        if not os.path.exists(resolved):
            # Try resolver fallback
            try:
                pr2 = PathResolver.resolve_user_file_any(file_path)
                if pr2 and pr2.exists():
                    resolved = str(pr2)
            except Exception:
                pass
        if not os.path.exists(resolved):
            raise FileNotFoundError(I18n.t("import.error.file_not_found", path=file_path))
        # Use resolved for next steps
        file_path = resolved
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext == '.docx':
            return self.import_from_docx(file_path)
        elif file_ext == '.doc':
            return self.import_from_doc(file_path)
        elif file_ext == '.pdf':
            return self.import_from_pdf(file_path)
        elif file_ext == '.xlsx':
            return self.import_from_xlsx(file_path)
        elif file_ext == '.txt':
            # Bo loc file da hua .txt nen phai xu ly that, khong bao loi
            # "dinh dang khong ho tro" khi nguoi dung chon dung file.
            return self.import_from_txt(file_path)
        else:
            raise ValueError(I18n.t("import.error.unsupported_format", ext=file_ext))

    def import_from_doc(self, file_path: str) -> List[Dict]:
        """Import questions from legacy DOC file (Windows COM fallback)"""
        # Try converting DOC -> DOCX via Word COM if available
        try:
            import win32com.client  # type: ignore
            import pythoncom  # type: ignore
        except ImportError:
            raise ImportError(I18n.t("import.error.doc_requires_win32com"))
        try:
            pythoncom.CoInitialize()
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(file_path)
            tmp_docx = str(Path(file_path).with_suffix('.converted.docx'))
            wdFormatXMLDocument = 12
            doc.SaveAs(tmp_docx, FileFormat=wdFormatXMLDocument)
            doc.Close(False)
            word.Quit()
            pythoncom.CoUninitialize()
            return self.import_from_docx(tmp_docx)
        except Exception as e:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass
            raise RuntimeError(I18n.t("import.error.doc_conversion_failed", error=str(e)))
    
    def import_from_docx(self, file_path: str) -> List[Dict]:
        """Import questions from DOCX file"""
        if not DOCX_AVAILABLE:
            raise ImportError(I18n.t("import.error.docx_missing"))
        file_path = self._resolve_user_file(file_path)
        questions = []
        
        try:
            # Handle OneDrive Files On-Demand hydration
            try:
                rp = PathResolver.resolve_user_file(file_path)
                if rp:
                    file_path = str(rp)
            except Exception:
                pass
            doc = _lazy_docx_document()(file_path)
            full_text = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                full_text.append(text)
            
            # Join text and parse questions
            text_content = '\n'.join(full_text)
            raw_questions = self._parse_questions_from_text(text_content)
            
            # Validate and normalize questions
            for q in raw_questions:
                validated = self._validate_imported_question(q)
                if validated:
                    questions.append(validated)
            
        except Exception as e:
            raise RuntimeError(I18n.t("import.error.docx_read_failed", error=str(e)))
        
        return questions
    
    def import_from_pdf(self, file_path: str) -> List[Dict]:
        """Import questions from PDF file"""
        if not PDF_AVAILABLE:
            raise ImportError(I18n.t("import.error.pdf_missing"))
        file_path = self._resolve_user_file(file_path)
        questions = []
        
        try:
            try:
                rp = PathResolver.resolve_user_file(file_path)
                if rp:
                    file_path = str(rp)
            except Exception:
                pass
            doc = _lazy_fitz().open(file_path)
            full_text = []
            
            # Extract text from all pages
            for page_num in range(doc.page_count):
                page = doc.load_page(page_num)
                text = page.get_text()
                if text.strip():
                    full_text.append(text.strip())
            
            doc.close()
            
            # Join text and parse questions
            text_content = '\n'.join(full_text)
            raw_questions = self._parse_questions_from_text(text_content)
            
            # Validate and normalize questions
            for q in raw_questions:
                validated = self._validate_imported_question(q)
                if validated:
                    questions.append(validated)
            
        except Exception as e:
            raise RuntimeError(I18n.t("import.error.pdf_read_failed", error=str(e)))
        
        return questions
    
    @staticmethod
    def _xlsx_cell_str(value) -> str:
        """Convert an Excel cell value to clean text (numbers stay clean: 100, not 100.0)."""
        if value is None:
            return ""
        if isinstance(value, bool):
            return "True" if value else "False"
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            return str(int(value)) if value.is_integer() else str(value)
        return str(value).strip()

    @staticmethod
    def _xlsx_norm(text: str) -> str:
        # Lowercase + bỏ dấu (kể cả đ→d vì NFD không tách được đ) để so khớp header/kiểu
        try:
            import unicodedata
            t = unicodedata.normalize("NFD", str(text or "").lower().replace("đ", "d"))
            t = "".join(c for c in t if unicodedata.category(c) != "Mn")
        except Exception:
            t = str(text or "").lower().replace("đ", "d")
        t = re.sub(r"[^a-z0-9\s]", " ", t)
        return re.sub(r"\s+", " ", t).strip()

    @staticmethod
    def _bracket_is_math(content: str) -> bool:
        """True when a [...] content is a math expression, not a blank.

        Rules: empty [] is always a blank; a leading +/- sign is ignored
        (so [-5] stays an answer); otherwise any operator (+-*/^=<>) means
        math, e.g. [x+5(78y+y^2)]. Plain words/numbers ([100], [Đại số])
        are answers.
        """
        try:
            s = (content or "").strip()
        except Exception:
            return False
        if not s:
            return False
        if s[0] in "+-" and len(s) > 1:
            s = s[1:].lstrip()
            if not s:
                return False
        return any(ch in _MATH_OPS for ch in s)

    @staticmethod
    def _space_run(text: str, idx: int, step: int) -> int:
        """Count consecutive spaces/tabs from idx (step -1 backward, +1 forward)."""
        try:
            n, j = 0, idx
            while 0 <= j < len(text) and text[j] in " \t":
                n += 1
                j += step
            return n
        except Exception:
            return 0

    def _split_brackets(self, text, known_answers=None):
        """Split [...] occurrences into real blanks vs literal math.

        Returns (display_text, blanks).
        Precedence:
        1. Double brackets [[...]] ALWAYS force a blank, even for math
           ("Biểu thức [[x+5(78y+y^2)]]+5 ..." -> blank "x+5...").
        2. Single [...] with known answers: a bracket matching a known
           answer is a blank; empty [] is always blank; non-math is blank;
           math brackets that match nothing stay literal.
        3. Single [...] with no known answers: an emphasized bracket
           (detached + 2+ spaces on either side) wins even if math; else
           every detached non-math bracket is a blank; if none is detached,
           all non-math contents are blanks (math stays literal).
           So "[x+1]+5" stays literal, "[100]" becomes a blank.
        known_answers cross-checks (e.g. Excel answer columns, '==' lines).
        """
        try:
            combined = list(_COMBINED_RE.finditer(text or ""))
        except Exception:
            return text, []
        if not combined:
            return text, []
        try:
            known = {self._xlsx_norm(a) for a in (known_answers or [])
                     if str(a or "").strip()}
        except Exception:
            known = set()

        def _key(content):
            try:
                return self._xlsx_norm(content)
            except Exception:
                return ""

        # Separate singles (doubles are always forced blanks).
        singles = []  # list of dicts with cidx, match, content, key, detached, emphasized, is_math, empty
        for cidx, m in enumerate(combined):
            try:
                dbl = m.group(1)
            except Exception:
                dbl = None
            if dbl is not None:
                continue
            try:
                content = m.group(2)
                if content is None:
                    content = ""
            except Exception:
                content = ""
            lead = self._space_run(text, m.start() - 1, -1)
            trail = self._space_run(text, m.end(), 1)
            try:
                detached = (m.start() == 0 or (text[m.start() - 1] in " \t")) \
                    and (m.end() == len(text) or (text[m.end()] in " \t"))
            except Exception:
                detached = False
            emphasized = lead >= 2 or trail >= 2
            try:
                is_math = self._bracket_is_math(content)
            except Exception:
                is_math = False
            try:
                empty = not str(content or "").strip()
            except Exception:
                empty = True
            singles.append({
                "cidx": cidx,
                "match": m,
                "content": content,
                "key": _key(content),
                "detached": detached,
                "emphasized": emphasized,
                "is_math": is_math,
                "empty": empty,
            })

        if known:
            # Rule 2: cross-check mode (singles only; doubles always blank).
            blanks_single_cidx = {
                inf["cidx"] for inf in singles
                if inf["empty"]
                or (inf["key"] and inf["key"] in known)
                or (not inf["is_math"])
            }
        else:
            # Rule 3: spacing decides; math stays literal unless emphasized.
            emph = {inf["cidx"] for inf in singles
                    if inf["detached"] and inf["emphasized"]}
            if emph:
                blanks_single_cidx = set(emph)
            else:
                det = {inf["cidx"] for inf in singles
                       if inf["detached"] and not inf["is_math"]}
                if det:
                    blanks_single_cidx = det
                else:
                    blanks_single_cidx = {inf["cidx"] for inf in singles
                                          if not inf["is_math"]}

        blanks = []
        parts, last = [], 0
        for cidx, m in enumerate(combined):
            try:
                dbl = m.group(1)
            except Exception:
                dbl = None
            if dbl is not None:
                parts.append(text[last:m.start()])
                parts.append("_____")
                last = m.end()
                blanks.append((dbl or "").strip())
            elif cidx in blanks_single_cidx:
                parts.append(text[last:m.start()])
                parts.append("_____")
                last = m.end()
                try:
                    blanks.append((m.group(2) or "").strip())
                except Exception:
                    blanks.append("")
            # else: literal math bracket stays in text via trailing slice
        parts.append(text[last:])
        return "".join(parts), blanks

    @staticmethod
    def _xlsx_formula_to_text(value) -> str:
        """Recover text from a formula-stored cell ('==' bị công cụ ghi thành công thức)."""
        s = str(value or "").strip()
        if len(s) > 1 and s[0] == "'" and s[1] == "=":
            s = s[1:]
        if s.startswith("=") and not s.startswith("=="):
            s = "=" + s
        return s.strip()

    def import_from_xlsx(self, file_path: str) -> List[Dict]:
        """Import questions from XLSX file.

        New template layout (STT | Câu hỏi | Đáp án 1..N | Giải thích,
        correct marked with '==' prefix) is detected by header; the legacy
        layout (Question | Option A-D | Correct | Explanation | Type) still works.
        """
        if not XLSX_AVAILABLE:
            raise ImportError(I18n.t("import.error.xlsx_missing"))
        file_path = self._resolve_user_file(file_path)
        questions = []

        try:
            try:
                rp = PathResolver.resolve_user_file(file_path)
                if rp:
                    file_path = str(rp)
            except Exception:
                pass
            wb = _lazy_load_workbook()(file_path, data_only=True)
            ws = wb.active
            # Workbook thứ 2 (giữ công thức dạng text) để cứu các ô '==' bị lưu thành công thức
            try:
                _wb_f = _lazy_load_workbook()(file_path, data_only=False)
                _ws_f = _wb_f.active
            except Exception:
                _wb_f = None
                _ws_f = None

            def _cell_text(r_idx: int, c_idx: int, cell) -> str:
                v = cell.value
                if v is None and _ws_f is not None:
                    try:
                        v = self._xlsx_formula_to_text(_ws_f.cell(row=r_idx, column=c_idx).value)
                    except Exception:
                        v = None
                s = self._xlsx_cell_str(v)
                # Bỏ dấu nháy openpyxl/Excel tự thêm trước '='
                if len(s) > 1 and s[0] == "'" and s[1] == "=":
                    s = s[1:]
                return s.strip()

            rows = []
            for r_idx, row_cells in enumerate(ws.iter_rows(), 1):
                rows.append([_cell_text(r_idx, c_idx, cell)
                             for c_idx, cell in enumerate(row_cells, 1)])
            # Drop fully empty rows but keep original indices for header search
            non_empty = [(i, r) for i, r in enumerate(rows)
                         if any(str(c or "").strip() for c in (r or []))]
            if not non_empty:
                wb.close()
                return []

            def _is_header(row) -> bool:
                cells = [self._xlsx_norm(c) for c in (row or [])]
                return any(c in ("cau hoi", "question") for c in cells)

            def _is_legacy_header(row) -> bool:
                cells = [self._xlsx_norm(c) for c in (row or [])]
                joined = " ".join(cells)
                return ("tuy chon" in joined or "option" in joined) and "dap an" not in joined.replace("dap an ", "")

            header_idx = -1
            legacy = False
            for i, r in non_empty:
                if _is_header(r):
                    header_idx = i
                    legacy = _is_legacy_header(r)
                    break

            if header_idx >= 0 and not legacy:
                # ---- New template layout ----
                header = non_empty[[k for k, (i, _) in enumerate(non_empty) if i == header_idx][0]][1]
                norm_headers = [self._xlsx_norm(c) for c in header]
                try:
                    col_q = next(k for k, c in enumerate(norm_headers) if c in ("cau hoi", "question"))
                except StopIteration:
                    col_q = 1
                try:
                    col_e = next(k for k, c in enumerate(norm_headers)
                                 if c in ("giai thich", "explanation") and k > col_q)
                except StopIteration:
                    col_e = len(header) - 1
                ans_cols = [k for k in range(col_q + 1, col_e)]
                if not ans_cols:
                    ans_cols = list(range(col_q + 1, len(header)))
                data_rows = [r for i, r in non_empty if i > header_idx]
                for r in data_rows:
                    cells = [self._xlsx_cell_str(c) for c in (r or [])]
                    while len(cells) <= max(col_q, col_e, max(ans_cols)):
                        cells.append("")
                    q_text = cells[col_q].strip()
                    if not q_text:
                        continue
                    if self._xlsx_norm(q_text) in ("cau hoi", "question", "stt",
                                                   "huong dan", "bang mau"):
                        continue
                    answers = [cells[k].strip() for k in ans_cols if cells[k].strip()]
                    if not answers:
                        continue
                    expl = cells[col_e].strip()
                    parsed = self._parse_xlsx_template_row(q_text, answers, expl)
                    if parsed is None:
                        continue
                    validated = self._validate_imported_question(parsed)
                    if validated:
                        questions.append(validated)
            else:
                # ---- Legacy layout: A=Question, B-E=Options, F=Correct, G=Explanation, H=Type ----
                start = header_idx + 1 if header_idx >= 0 else 1
                for r in rows[start:]:
                    if not r or not str(r[0] or "").strip():
                        continue
                    question_data = {
                        "question": str(r[0]) if r[0] else "",
                        "options": [],
                        "correct_answer": str(r[5]) if len(r) > 5 and r[5] else "",
                        "explanation": str(r[6]) if len(r) > 6 and r[6] else "",
                        "type": str(r[7]).lower().replace(" ", "_") if len(r) > 7 and r[7] else "multiple_choice"
                    }
                    if len(r) > 4:
                        options = []
                        for i, option in enumerate(r[1:5], 1):
                            if option:
                                options.append(f"{chr(64 + i)}. {option}")
                        question_data["options"] = options
                    validated_question = self._validate_imported_question(question_data)
                    if validated_question:
                        questions.append(validated_question)

            try:
                wb.close()
            except Exception:
                pass
            try:
                if _wb_f is not None:
                    _wb_f.close()
            except Exception:
                pass

        except Exception as e:
            raise RuntimeError(I18n.t("import.error.xlsx_read_failed", error=str(e)))

        return questions

    def _parse_xlsx_template_row(self, q_text: str, answers: List[str], expl: str) -> Optional[Dict]:
        """Parse one new-template Excel row (STT | Câu hỏi | Đáp án.. | Giải thích).

        Conventions: '==' prefix marks the correct choice, 'left = right'
        cells are matching pairs, a question containing [] or ___ is fill blank.
        Leading 'Câu hỏi N:' prefixes are stripped.
        """
        q_text = re.sub(
            r"^(câu(\s+hỏi)?|question|pregunta|frage)\s+\d+\s*[:\.]?\s*",
            "", q_text.strip(), flags=re.IGNORECASE)
        # Strip '==' correct markers (keep position) and letter prefixes (A. / B) ...)
        cleaned = []
        correct_idx = -1
        for ans in answers:
            text = ans.strip()
            if text.startswith("=="):
                text = text[2:].strip()
                if correct_idx == -1:
                    correct_idx = len(cleaned)
            text = re.sub(r"^([A-Fa-f]|\d{1,2})[\.\)]\s+", "", text).strip()
            if text:
                cleaned.append(text)
        _has_blanks = ("[" in q_text and "]" in q_text) or ("___" in q_text)
        if not cleaned and not _has_blanks:
            return None

        # Fill in the blank: real blanks become ___ but math brackets
        # (e.g. [x+5(78y+y^2)]) stay literal in the question.
        display_q, blanks = self._split_brackets(q_text, cleaned)
        has_us = "___" in q_text
        if blanks or has_us:
            fb_answers = list(cleaned) if cleaned else [b for b in blanks if b]
            if not fb_answers:
                return None
            return {
                "type": "fill_blank",
                "question": display_q,
                "correct_answers": fb_answers,
                "answers": fb_answers,
                "correct_answer": fb_answers[0],
                "explanation": expl,
            }

        # Matching: every answer cell is a 'left = right' pair
        def _as_pair(text: str):
            if any(op in text for op in (">=", "<=", "=>", "==")):
                return None
            if "=" not in text:
                return None
            left, right = text.split("=", 1)
            left, right = left.strip(), right.strip()
            return {"left": left, "right": right} if left and right else None

        pairs = [_as_pair(t) for t in cleaned]
        if len(cleaned) >= 2 and all(pairs):
            return {
                "type": "matching",
                "question": q_text,
                "pairs": pairs,
                "explanation": expl,
            }

        # True/False: the two answers are a Đúng/Sai pair
        tf_true = {"đúng", "dung", "true", "vrai", "verdadero", "wahr"}
        tf_false = {"sai", "false", "faux", "falso", "falsch"}
        lowered = [self._xlsx_norm(t) for t in cleaned]
        if len(cleaned) == 2 and (
            (lowered[0] in tf_true and lowered[1] in tf_false)
            or (lowered[0] in tf_false and lowered[1] in tf_true)
        ):
            if correct_idx >= 0:
                correct_bool = cleaned[correct_idx].strip().lower() not in (
                    "sai", "false", "faux", "falso", "falsch")
            else:
                correct_bool = lowered[0] in tf_true
            return {
                "type": "true_false",
                "question": q_text,
                "correct_answer": correct_bool,
                "explanation": expl,
            }

        # Short answer: single answer, or several case variants of one answer
        # without any '==' marker (e.g. "William Shakespeare" / "william shakespeare")
        if len(cleaned) == 1:
            return {
                "type": "short_answer",
                "question": q_text,
                "answers": cleaned,
                "explanation": expl,
            }
        if correct_idx == -1:
            try:
                _normed = [self._xlsx_norm(t) for t in cleaned]
            except Exception:
                _normed = list(cleaned)
            if len(set(_normed)) == 1:
                return {
                    "type": "short_answer",
                    "question": q_text,
                    "answers": cleaned,
                    "explanation": expl,
                }

        # Default: multiple choice
        return {
            "type": "multiple_choice",
            "question": q_text,
            "options": cleaned,
            "correct_answer": correct_idx if correct_idx >= 0 else 0,
            "explanation": expl,
        }
    
    def _parse_questions_from_text(self, text: str) -> List[Dict]:
        """Parse questions from plain text"""
        questions = []
        
        try:
            text = self._normalize_input(text)
        except Exception:
            text = text.replace('\r\n', '\n').strip()
        
        # Parse using specific parsers
        # Try smart parse first as it handles blocks better
        smart_questions = self.parse_smart_format(text)
        if smart_questions:
            questions.extend(smart_questions)
        
        # Legacy parsers are disabled to prevent duplicates and ensure consistent explanation extraction
        # questions.extend(self.parse_multiple_choice_questions(text))
        # questions.extend(self.parse_true_false_questions(text))
        # questions.extend(self.parse_fill_blank_questions(text))
        # questions.extend(self.parse_matching_questions(text))
        
        return questions

    def _parse_custom_marker_blocks(self, text: str) -> List[Dict]:
        """Parse questions using custom markers:
        * starts a question
        Options are lines that look like A., a), 1), a., 1) ...
        Correct options are prefixed with '=='
        Explanation starts with '/' and continues until next '*'
        Supports multiple_choice, true_false, fill_blank, matching (left - right)
        """
        lines = [l.rstrip() for l in text.splitlines()]
        blocks: List[List[str]] = []
        current: Optional[List[str]] = None
        for ln in lines:
            stripped = ln.strip()
            if stripped.startswith('*'):
                if current:
                    blocks.append(current)
                current = [stripped[1:].strip()]
                continue
            if current is None:
                continue
            current.append(ln)
        if current:
            blocks.append(current)
        results: List[Dict] = []
        for blk in blocks:
            if not blk:
                continue
            question_text_parts: List[str] = []
            options: List[Dict] = []
            pairs: List[Dict] = []
            explanation_parts: List[str] = []
            in_explain = False
            seen_option = False
            for idx, raw in enumerate(blk):
                raw_stripped = str(raw).strip()
                if in_explain and re.match(r'^[A-D]\.\s+.+', raw_stripped):
                    break
                if raw_stripped.startswith('/'):
                    in_explain = True
                    explanation_parts.append(raw_stripped[1:].strip())
                    continue
                if in_explain:
                    explanation_parts.append(raw_stripped)
                    continue
                # Matching pair: left - right
                m_pair = re.match(r'^\s*(==)?\s*(.+?)\s*[-–>→]\s*(.+)\s*$', raw_stripped)
                if m_pair and not seen_option:
                    pairs.append({"left": m_pair.group(2).strip(), "right": m_pair.group(3).strip()})
                    continue
                # Option markers
                m_opt = re.match(r'^\s*(==)?\s*([A-Da-d]|\d{1,2})[\.\)]\s*(.+)\s*$', raw_stripped)
                if m_opt:
                    seen_option = True
                    is_correct = bool(m_opt.group(1))
                    opt_text = m_opt.group(3).strip()
                    options.append({"text": opt_text, "correct": is_correct})
                    continue
                # Continuation lines
                if seen_option and raw_stripped:
                    try:
                        options[-1]["text"] = f"{options[-1]['text']} {raw_stripped}".strip()
                    except Exception:
                        pass
                    continue
                
                # Ignore answer lines for fill_blank (starting with == but not option/pair)
                if raw_stripped.startswith('==') and idx > 0:
                    continue

                # Question text parts
                if raw_stripped:
                    question_text_parts.append(raw_stripped)
            q_text = ' '.join([p for p in question_text_parts if p]).strip()
            expl = ' '.join([p for p in explanation_parts if p]).strip()
            if pairs and not options:
                results.append({
                    "type": "matching",
                    "question": q_text or I18n.t("import.default.matching_question"),
                    "pairs": pairs,
                    "explanation": expl
                })
                continue
            if not options:
                # Fill blank: use any line prefixed with '==' as answer otherwise leave empty
                answer_lines = [re.sub(r'^\s*==\s*', '', l).strip() for l in blk if str(l).strip().startswith('==')]
                results.append({
                    "type": "fill_blank",
                    "question": q_text or I18n.t("import.default.fill_blank_question"),
                    "answers": answer_lines if answer_lines else [],
                    "case_sensitive": False,
                    "explanation": expl
                })
                continue
            # Determine type true/false if 2 options and match
            texts_lower = [str(o.get('text','')).strip().lower() for o in options]
            if len(options) == 2 and set(texts_lower) <= {"true","false","đúng","sai"}:
                correct_bool = options[0].get('correct', False)
                # if no marker, default first to True
                if not any(o.get('correct') for o in options):
                    correct_bool = texts_lower[0] in {"true","đúng"}
                results.append({
                    "type": "true_false",
                    "question": q_text,
                    "correct_answer": True if correct_bool else False,
                    "explanation": expl
                })
                continue
            # Multiple choice: build A-D options text list and correct letter
            texts = [o['text'] for o in options]
            letters = ['A','B','C','D']
            correct_letter = None
            for idx, o in enumerate(options):
                if o.get('correct'):
                    try:
                        correct_letter = letters[idx]
                    except Exception:
                        correct_letter = None
                    break
            results.append({
                "type": "multiple_choice",
                "question": q_text,
                "options": [f"{letters[i] if i < len(letters) else str(i+1)}. {texts[i]}" for i in range(len(texts))],
                "correct_answer": correct_letter or 'A',
                "explanation": expl
            })
        return results
    
    def _extract_question_from_match(self, match) -> Optional[Dict]:
        """Extract question data from regex match"""
        groups = match.groups()
        
        if len(groups) < 2:
            return None
        
        question_data = {
            "question": groups[1].strip(),
            "options": [],
            "correct_answer": "",
            "explanation": "",
            "type": "multiple_choice"
        }
        
        # Extract options (A, B, C, D)
        if len(groups) >= 6:
            for i in range(2, 6):
                if i < len(groups) and groups[i]:
                    option_text = groups[i].strip()
                    if option_text and not option_text.startswith('Answer:'):
                        question_data["options"].append(option_text)
        
        # Extract answer
        if len(groups) >= 7 and groups[6]:
            answer_text = groups[6].strip().upper()
            if len(answer_text) == 1 and answer_text in 'ABCD':
                question_data["correct_answer"] = answer_text
        
        # Extract explanation
        if len(groups) >= 8 and groups[7]:
            explanation = groups[7].strip()
            if explanation and not explanation.startswith('['):
                question_data["explanation"] = explanation
        
        return question_data
    
    def _validate_imported_question(self, question_data: Dict) -> Optional[Dict]:
        """Validate and fix imported question data"""
        if not question_data.get("question"):
            return None
        
        # Ensure required fields
        question_data["type"] = question_data.get("type", "multiple_choice")
        question_data["explanation"] = question_data.get("explanation", "")
        
        qt = (question_data.get("type") or "multiple_choice").lower().strip()
        
        # Map common type names to internal names
        if qt in ("fill_blank", "fill_in_the_blank", "dien_khuyet"):
            qt = "fill_blank"
        elif qt in ("matching", "ghep_cap"):
            qt = "matching"
        elif qt in ("true_false", "dung_sai"):
            qt = "true_false"
        elif qt in ("multiple_choice", "trac_nghiem"):
            qt = "multiple_choice"
            
        question_data["type"] = qt

        if qt == "multiple_choice":
            # Handle list of dicts for options (from smart parser)
            opts = question_data.get("options")
            if opts and isinstance(opts, list) and len(opts) > 0 and isinstance(opts[0], dict):
                # Extract text and find correct answer
                new_options = []
                correct_idx = -1
                for i, opt in enumerate(opts):
                    text = opt.get("text", "")
                    new_options.append(text)
                    if opt.get("correct"):
                        correct_idx = i
                
                question_data["options"] = new_options
                # If correct_answer wasn't set explicitly, use the one found in options
                if question_data.get("correct_answer") is None or question_data.get("correct_answer") == "":
                    if correct_idx != -1:
                        question_data["correct_answer"] = correct_idx

            if not question_data.get("options"):
                question_data["options"] = [
                    f"A. {I18n.t('import.default.option_1')}",
                    f"B. {I18n.t('import.default.option_2')}", 
                    f"C. {I18n.t('import.default.option_3')}",
                    f"D. {I18n.t('import.default.option_4')}"
                ]
            
            # Normalize correct_answer to index (int)
            ca = question_data.get("correct_answer")
            if ca is None or ca == "":
                question_data["correct_answer"] = 0
            elif isinstance(ca, str):
                # Convert "A", "B", "C", "D" to 0, 1, 2, 3
                clean_ca = ca.strip().upper()
                if len(clean_ca) == 1 and clean_ca in "ABCDEF":
                    question_data["correct_answer"] = ord(clean_ca) - ord('A')
                elif clean_ca.isdigit():
                    question_data["correct_answer"] = int(clean_ca)
            # If it's already an int (like 0), leave it alone

                
        elif qt == "true_false":
            # Ensure correct_answer is boolean or A/B
            ca = question_data.get("correct_answer")
            if isinstance(ca, str):
                ca = ca.strip().upper()
                if ca in ('A', 'TRUE', 'ĐÚNG', 'DUNG'):
                    question_data["correct_answer"] = True
                else:
                    question_data["correct_answer"] = False
            elif ca is None:
                 question_data["correct_answer"] = True
                 
        elif qt == "fill_blank":
            # Ensure correct_answers is a list
            ans = question_data.get("answers") or question_data.get("correct_answers")
            if not ans:
                ans = []
            if not isinstance(ans, list):
                ans = [str(ans)]
            question_data["correct_answers"] = ans
            # Remove legacy field if present
            if "answers" in question_data:
                del question_data["answers"]
                
        elif qt == "matching":
            # Ensure pairs is a list of dicts
            pairs = question_data.get("pairs") or []
            if not isinstance(pairs, list):
                pairs = []
            # Validate pair structure
            valid_pairs = []
            for p in pairs:
                if isinstance(p, dict) and 'left' in p and 'right' in p:
                    valid_pairs.append(p)
            question_data["pairs"] = valid_pairs
        
        # Clean up question text
        question_data["question"] = self._clean_text(question_data["question"])
        question_data["explanation"] = self._clean_text(question_data["explanation"])
        
        # Clean up options
        if "options" in question_data:
            question_data["options"] = [self._clean_text(opt) for opt in question_data["options"] if opt.strip()]
        
        return question_data
    
    def _clean_text(self, text: str) -> str:
        """Clean and normalize text"""
        if not text:
            return ""

        try:
            text = self._normalize_input(text)
        except Exception:
            pass

        text = re.sub(r'\s+', ' ', str(text).strip())
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        return text.strip()
    
    def create_sample_template(self, file_path: str, format_type: str = "docx", lang: str = "vi", template_type: str = "general"):
        """Create a sample template file for teachers"""
        if format_type == "docx" and DOCX_AVAILABLE:
            self._create_sample_docx(file_path, lang, template_type)
        elif format_type == "doc":
            self._create_sample_doc(file_path, lang, template_type)
        elif format_type == "xlsx" and XLSX_AVAILABLE:
            self._create_sample_xlsx(file_path, lang, template_type)
        elif format_type == "txt":
            self._create_sample_txt(file_path, lang, template_type)
        else:
            raise ValueError(I18n.t("import.error.template_creation_not_supported", format=format_type))

    def _create_sample_txt(self, file_path: str, lang: str, template_type: str = "general"):
        """Create sample TXT template by copying from assets"""
        # Determine source file name
        is_vi = lang == "vi"
        
        if template_type == "millionaire":
            src_name = "mau_cau_hoi_trieu_phu_vi.txt" if is_vi else "sample_questions_millionaire_en.txt"
        else:
            src_name = "mau_cau_hoi_tong_hop_vi.txt" if is_vi else "sample_questions_general_en.txt"
            
        src_path = self.templates_dir / src_name
        
        if not src_path.exists():
            # Fallback to English general if specific not found
            src_path = self.templates_dir / "sample_questions_general_en.txt"
            
        if not src_path.exists():
             # Last resort fallback content
             with open(file_path, "w", encoding="utf-8") as f:
                 f.write("Template file not found.")
             return

        # Copy content
        import shutil
        shutil.copy2(src_path, file_path)
    
    def _create_sample_docx(self, file_path: str, lang: str, template_type: str = "general"):
        """Create sample DOCX template"""
        doc = _lazy_docx_document()()
        l = lang or "vi"
        
        doc.add_heading(I18n.t("template.docx.title", l), 0)
        doc.add_paragraph(I18n.t("template.docx.subtitle", l))
        
        # Add notes about game types
        if template_type == "millionaire":
             doc.add_paragraph(I18n.t("template.docx.note_millionaire", l))
        else:
             doc.add_paragraph(I18n.t("template.docx.note_millionaire", l))
             doc.add_paragraph(I18n.t("template.docx.note_others", l))
        doc.add_paragraph("")
        
        doc.add_heading(I18n.t("template.docx.instructions_heading", l), level=1)
        doc.add_paragraph(I18n.t("template.docx.instructions_1", l))
        doc.add_paragraph(I18n.t("template.docx.instructions_2", l))
        doc.add_paragraph(I18n.t("template.docx.instructions_3", l))
        doc.add_paragraph(I18n.t("template.docx.instructions_4", l))
        # ALTP/millionaire: shortened type list (MC + True/False only).
        if str(template_type or "").lower() == "millionaire":
            _i5 = I18n.t("template.docx.instructions_5_millionaire", l)
            doc.add_paragraph(_i5 if _i5 and _i5 != "template.docx.instructions_5_millionaire"
                             else I18n.t("template.docx.instructions_5", l))
        else:
            doc.add_paragraph(I18n.t("template.docx.instructions_5", l))
        doc.add_paragraph("")

        # Separate attention block about [...] vs [[...]] (not merged into instructions)
        # ALTP/millionaire: no attention needed (MC only).
        if str(template_type or "").lower() != "millionaire":
            try:
                _att_title = I18n.t("template.docx.attention_title", l)
                _att_1 = I18n.t("template.docx.attention_1", l)
                _att_2 = I18n.t("template.docx.attention_2", l)
                _att_3 = I18n.t("template.docx.attention_3", l)
            except Exception:
                _att_title, _att_1, _att_2, _att_3 = "", "", "", ""
            if _att_title and _att_title != "template.docx.attention_title":
                try:
                    from docx.shared import RGBColor
                    _h = doc.add_heading(_att_title, level=2)
                    try:
                        for _run in _h.runs:
                            _run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                    except Exception:
                        pass
                except Exception:
                    doc.add_heading(_att_title, level=2)
                for _key in (_att_1, _att_2, _att_3):
                    try:
                        if _key and not str(_key).startswith("template.docx.attention_"):
                            _p = doc.add_paragraph(str(_key))
                            try:
                                from docx.shared import Pt
                                for _run in _p.runs:
                                    _run.font.size = Pt(10.5)
                            except Exception:
                                pass
                    except Exception:
                        pass
                doc.add_paragraph("")
        
        doc.add_heading(I18n.t("template.docx.samples_heading", l), level=1)
        doc.add_paragraph("")

        # 1. Multiple Choice (Always included)
        doc.add_heading(I18n.t("template.docx.sample_mc_heading", l), level=2)
        doc.add_paragraph(I18n.t("template.docx.sample_mc_q", l))
        doc.add_paragraph(I18n.t("template.docx.sample_mc_a1", l))
        doc.add_paragraph(I18n.t("template.docx.sample_mc_a2", l))
        doc.add_paragraph(I18n.t("template.docx.sample_mc_a3", l))
        doc.add_paragraph(I18n.t("template.docx.sample_mc_a4", l))
        doc.add_paragraph(I18n.t("template.docx.sample_mc_expl", l))
        doc.add_paragraph("")

        if template_type != "millionaire":
            # 2. True/False
            doc.add_heading(I18n.t("template.docx.sample_tf_heading", l), level=2)
            doc.add_paragraph(I18n.t("template.docx.sample_tf_q", l))
            doc.add_paragraph(I18n.t("template.docx.sample_tf_true", l))
            doc.add_paragraph(I18n.t("template.docx.sample_tf_false", l))
            doc.add_paragraph(I18n.t("template.docx.sample_tf_expl", l))
            doc.add_paragraph("")

            # 3. Fill in the Blank
            doc.add_heading(I18n.t("template.docx.sample_fb_heading", l), level=2)
            doc.add_paragraph(I18n.t("template.docx.sample_fb_q", l))
            fb_ans = I18n.t("template.docx.sample_fb_answer", l)
            if fb_ans:
                doc.add_paragraph(fb_ans)
            doc.add_paragraph(I18n.t("template.docx.sample_fb_expl", l))
            doc.add_paragraph("")

            # 4. Matching
            doc.add_heading(I18n.t("template.docx.sample_match_heading", l), level=2)
            doc.add_paragraph(I18n.t("template.docx.sample_match_q", l))
            doc.add_paragraph(I18n.t("template.docx.sample_match_pair1", l))
            doc.add_paragraph(I18n.t("template.docx.sample_match_pair2", l))
            doc.add_paragraph(I18n.t("template.docx.sample_match_expl", l))
            doc.add_paragraph("")

            # 5. Short Answer
            doc.add_heading(I18n.t("template.docx.sample_sa_heading", l), level=2)
            doc.add_paragraph(I18n.t("template.docx.sample_sa_q", l))
            doc.add_paragraph(I18n.t("template.docx.sample_sa_answer", l))
            doc.add_paragraph(I18n.t("template.docx.sample_sa_expl", l))
            doc.add_paragraph("")
        
        doc.save(file_path)
    
    def _create_sample_xlsx(self, file_path: str, lang: str = "vi", template_type: str = "general"):
        """Create sample XLSX template (guide block + yellow sample table, Times New Roman)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill
        from openpyxl.utils import get_column_letter

        l = lang or "vi"
        wb = Workbook()
        ws = wb.active
        ws.title = I18n.t("template.xlsx.sheet_title", l)
        try:
            ws.sheet_view.showGridLines = True
        except Exception:
            pass

        _TNR = "Times New Roman"
        _BLUE = "0000FF"
        _YELLOW = "FFFF00"
        _thin_side = __import__("openpyxl.styles", fromlist=["Side"]).Side(style="thin", color="000000")
        _thin_border = Border(left=_thin_side, right=_thin_side, top=_thin_side, bottom=_thin_side)

        def _t(key: str, **kwargs) -> str:
            try:
                val = I18n.t(key, l, **kwargs)
            except Exception:
                val = ""
            if not val or val == key:
                return ""
            return val

        # Row 1: guide title (blue, bold), merged A1:B1 like the reference file
        ws.merge_cells("A1:B1")
        ws["A1"] = _t("template.xlsx.guide_title") or "HƯỚNG DẪN:"
        ws["A1"].font = Font(name=_TNR, size=14, bold=True, color=_BLUE)

        # Row 2: spacer (merged A2:B2)
        ws.merge_cells("A2:B2")

        # Rows 3-4 merged: guide text (wrapped, top aligned)
        # ALTP/millionaire: only 2 items; general: 4 items. Note block is separate.
        _is_altp = str(template_type or "").lower() == "millionaire"
        _guide_keys = ("template.xlsx.guide_1", "template.xlsx.guide_2") if _is_altp else (
            "template.xlsx.guide_1",
            "template.xlsx.guide_2",
            "template.xlsx.guide_3",
            "template.xlsx.guide_4",
        )
        guide = "\n".join([x for x in (_t(_k) for _k in _guide_keys) if x])
        ws.merge_cells("A3:G4")
        _g = ws["A3"]
        _g.value = guide
        _g.font = Font(name=_TNR, size=12)
        _g.alignment = Alignment(wrap_text=True, vertical="top")
        try:
            ws.row_dimensions[3].height = 30 if _is_altp else 60
            ws.row_dimensions[4].height = 15
        except Exception:
            pass

        # Row 5: separate attention title (red, bold) - NOT merged into guide
        # ALTP/millionaire: no attention needed (MC only).
        _RED = "C00000"
        _show_note = not _is_altp
        _note_title = _t("template.xlsx.note_title") if _show_note else ""
        ws.merge_cells("A5:G5")
        _nt = ws["A5"]
        _nt.value = _note_title
        _nt.font = Font(name=_TNR, size=12, bold=True, color=_RED)
        _nt.alignment = Alignment(vertical="center", wrap_text=True)
        try:
            ws.row_dimensions[5].height = (18 if _note_title else 5)
        except Exception:
            pass

        # Rows 6-7 merged: attention body (nice, easy to read)
        _note_body = ("\n".join([x for x in (
            _t("template.xlsx.note_1"),
            _t("template.xlsx.note_2"),
            _t("template.xlsx.note_3"),
        ) if x]) if _show_note else "")
        ws.merge_cells("A6:G7")
        _nb = ws["A6"]
        _nb.value = _note_body
        _nb.font = Font(name=_TNR, size=11, color="404040")
        _nb.alignment = Alignment(wrap_text=True, vertical="top")
        try:
            ws.row_dimensions[6].height = (45 if _show_note else 5)
            ws.row_dimensions[7].height = (15 if _show_note else 5)
        except Exception:
            pass

        # Row 8: spacer (merged A8:G8)
        ws.merge_cells("A8:G8")

        # Row 9: sample table title (merged, centered)
        ws.merge_cells("A9:G9")
        _tt = ws["A9"]
        _tt.value = _t("template.xlsx.table_title") or "BẢNG MẪU"
        _tt.font = Font(name=_TNR, size=12, bold=False)
        _tt.alignment = Alignment(horizontal="center", vertical="center")
        try:
            ws.row_dimensions[9].height = 22
        except Exception:
            pass

        # Sample rows: (question, [answers...], explanation)
        samples = [
            (
                _t("template.xlsx.s1_q"),
                [_t("template.xlsx.s1_a1"), _t("template.xlsx.s1_a2"),
                 _t("template.xlsx.s1_a3"), _t("template.xlsx.s1_a4")],
                _t("template.xlsx.s1_e"),
            ),
            (
                _t("template.xlsx.s2_q"),
                [_t("template.xlsx.s2_a1"), _t("template.xlsx.s2_a2")],
                _t("template.xlsx.s2_e"),
            ),
            (
                _t("template.xlsx.s3_q"),
                [_t("template.xlsx.s3_a1")],
                _t("template.xlsx.s3_e"),
            ),
            (
                _t("template.xlsx.s4_q"),
                [_t("template.xlsx.s4_a1"), _t("template.xlsx.s4_a2")],
                _t("template.xlsx.s4_e"),
            ),
            (
                _t("template.xlsx.s5_q"),
                [_t("template.xlsx.s5_a1"), _t("template.xlsx.s5_a2")],
                _t("template.xlsx.s5_e"),
            ),
        ]
        if template_type == "millionaire":
            samples = samples[:2]

        _n_ans = 4
        headers = (
            [_t("template.xlsx.header_stt") or "STT",
             _t("template.xlsx.header_question") or "Câu hỏi"]
            + [_t("template.xlsx.header_answer", n=i) or f"Đáp án {i}" for i in range(1, _n_ans + 1)]
            + [_t("template.xlsx.header_explain") or "Giải thích"]
        )
        _last_col = len(headers)
        _last_letter = get_column_letter(_last_col)

        # Row 10: yellow header
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=10, column=col_idx, value=header)
            cell.font = Font(name=_TNR, size=12, bold=True)
            cell.fill = PatternFill(start_color=_YELLOW, end_color=_YELLOW, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _thin_border
        try:
            ws.row_dimensions[10].height = 22
        except Exception:
            pass

        # Rows 11+: sample data with thin borders
        for r_offset, (q_text, answers, expl) in enumerate(samples):
            row_idx = 11 + r_offset
            padded = list(answers[:_n_ans]) + [""] * max(0, _n_ans - len(answers))
            values = [r_offset + 1, q_text] + padded + [expl]
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                # Force inline text so leading '==' is NOT stored as a formula
                try:
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"
                except Exception:
                    pass
                cell.font = Font(name=_TNR, size=12)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = _thin_border

        # Fixed column widths like the reference layout
        for _letter, _width in (("A", 3.9), ("B", 30.9), ("C", 21.9), ("D", 24.7),
                                ("E", 18.4), ("F", 16.4), ("G", 47.5)):
            try:
                ws.column_dimensions[_letter].width = _width
            except Exception:
                pass

        wb.save(file_path)

    def _create_sample_doc(self, file_path: str, lang: str = "vi", template_type: str = "general"):
        """Create sample DOC template by generating DOCX then converting via Word COM if available"""
        try:
            tmp_docx = str(Path(file_path).with_suffix('.tmp.docx'))
            self._create_sample_docx(tmp_docx, lang, template_type)
            try:
                import win32com.client  # type: ignore
                import pythoncom  # type: ignore
            except ImportError:
                # Fallback: just rename to .docx if COM not available
                Path(tmp_docx).replace(str(Path(file_path).with_suffix('.docx')))
                return
            try:
                pythoncom.CoInitialize()
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(tmp_docx)
                wdFormatDocument = 0
                doc.SaveAs(file_path, FileFormat=wdFormatDocument)
                doc.Close(False)
                word.Quit()
                pythoncom.CoUninitialize()
            finally:
                try:
                    Path(tmp_docx).unlink(missing_ok=True)
                except Exception:
                    pass
        except Exception as e:
            raise RuntimeError(I18n.t("import.error.doc_template_failed", error=str(e)))

    def split_into_chapters(self, file_path: str):
        """Tách tài liệu thành các chương/bài tự động (hiểu sâu) - dùng cho AI tạo theo chương"""
        file_path = self._resolve_user_file(file_path)
        text = ""
        try:
            ext = Path(file_path).suffix.lower()
            if ext == ".txt":
                text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
            elif ext == ".pdf" and PDF_AVAILABLE:
                doc = _lazy_fitz().open(file_path)
                parts = []
                for i in range(min(10, doc.page_count)):
                    parts.append(doc.load_page(i).get_text("text"))
                text = "\n".join(parts)
            elif ext in [".docx", ".doc"] and DOCX_AVAILABLE:
                doc = _lazy_docx_document()(file_path)
                text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            else:
                text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        except Exception:
            text = ""
        if not text.strip():
            return []
        import re
        pattern = r'(?=^(?:Chương|Bài|Unit|Chapter|CHƯƠNG|BÀI)\s+[0-9IVXLC]+\b.*$)'
        parts = re.split(pattern, text, flags=re.MULTILINE | re.IGNORECASE)
        parts = [p.strip() for p in parts if p.strip()]
        chapters = []
        for idx, part in enumerate(parts, 1):
            first_line = part.splitlines()[0].strip()[:80] if part.splitlines() else f"Chương {idx}"
            title = first_line if re.match(r'^(?:Chương|Bài|Unit|Chapter)', first_line, re.IGNORECASE) else f"Phần {idx}: {first_line[:40]}"
            chapters.append({"index": idx, "title": title, "content": part[:3000], "char_count": len(part)})
        return chapters

    def enhanced_dedup_check(self, new_question, existing):
        """Kiểm tra trùng nâng cao: so sánh cả đáp án và ngữ nghĩa (dùng cho AI hiểu sâu)"""
        try:
            from eduplay.core.project_manager import ProjectManager
            pm = ProjectManager()
            base_dup = pm._is_duplicate_question(new_question, existing)
            if base_dup:
                return True
            import re, unicodedata
            def _norm(s):
                s = str(s or "").lower()
                s = unicodedata.normalize("NFD", s)
                s = "".join(c for c in s if unicodedata.category(c) != "Mn")
                s = re.sub(r'[^a-z0-9\s]', ' ', s)
                s = re.sub(r'\s+', ' ', s).strip()
                return s
            new_opts = new_question.get("options") or []
            new_q = _norm(new_question.get("question") or "")
            for ex in existing:
                ex_q = _norm(ex.get("question") or "")
                if new_q and ex_q:
                    set_new = set(new_q.split())
                    set_ex = set(ex_q.split())
                    if set_new and set_ex:
                        jacc = len(set_new & set_ex) / len(set_new | set_ex)
                        if jacc > 0.85:
                            ex_opts = ex.get("options") or []
                            if new_opts and ex_opts:
                                norm_new_opts = set(_norm(o) for o in new_opts if o)
                                norm_ex_opts = set(_norm(o) for o in ex_opts if o)
                                if len(norm_new_opts & norm_ex_opts) / max(len(norm_new_opts), 1) > 0.8:
                                    return True
            return False
        except Exception:
            return False
"""
Nguyen-Thanh-Tan ¬_¬
"""
